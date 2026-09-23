"""
IRL 3 - Presentasi Konsep Dashboard (Simplified)
Fokus: Simple, Singkat, + Penjelasan Perhitungan
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ─────────────────────────────────────────────────
NAVY        = RGBColor(0x1E, 0x3A, 0x5F)
TEAL        = RGBColor(0x00, 0x7A, 0x87)
CYAN        = RGBColor(0x00, 0xBF, 0xD8)
GREEN       = RGBColor(0x2E, 0x7D, 0x32)
AMBER       = RGBColor(0xF5, 0x7C, 0x00)
RED         = RGBColor(0xC6, 0x28, 0x28)
GRAY        = RGBColor(0xF4, 0xF6, 0xF9)
GRAY_MID    = RGBColor(0x78, 0x90, 0x9C)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
DARK        = RGBColor(0x1A, 0x1A, 0x2E)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]
TOTAL = 6

def new_slide(): return prs.slides.add_slide(BL)

def R(sl, l, t, w, h, fill=WHITE, line=None, lw=1, radius=False):
    s = sl.shapes.add_shape(5 if radius else 1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line: s.line.color.rgb = line; s.line.width = Pt(lw)
    else: s.line.fill.background()
    return s

def T(sl, text, l, t, w, h, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT, italic=False):
    tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    r  = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = color
    return tb

def title_bar(sl, title, sub=None):
    R(sl, 0, 0, 13.33, 1.0, fill=NAVY)
    R(sl, 0, 0.9, 13.33, 0.1, fill=CYAN)
    T(sl, title, 0.4, 0.08, 12.5, 0.6, size=24, bold=True, color=WHITE)
    if sub: T(sl, sub, 0.4, 0.65, 12.5, 0.3, size=11, color=CYAN)

def footer(sl, num):
    R(sl, 0, 7.28, 13.33, 0.22, fill=NAVY)
    T(sl, "Dashboard GM — Material Ready", 0.2, 7.29, 11, 0.2, size=7.5, color=RGBColor(0xBB,0xCC,0xDD))
    T(sl, f"{num} / {TOTAL}", 12.5, 7.29, 0.7, 0.2, size=7.5, color=RGBColor(0x88,0x99,0xAA), align=PP_ALIGN.RIGHT)

def bullet_list(sl, items, l, t, w, size=11, gap=0.4, color=DARK, dot_color=CYAN):
    y = t
    for item in items:
        R(sl, l, y+0.09, 0.12, 0.12, fill=dot_color)
        T(sl, item, l+0.22, y, w-0.25, gap, size=size, color=color)
        y += gap
    return y


# ==========================================
# 1. COVER
# ==========================================
sl = new_slide()
R(sl, 0, 0, 13.33, 7.5, fill=NAVY)
T(sl, "IRL 3", 1.0, 2.0, 10, 0.6, size=16, bold=True, color=CYAN)
T(sl, "Konsep Dashboard Material Ready", 1.0, 2.5, 10, 1.3, size=40, bold=True, color=WHITE)
T(sl, "Diskusi dengan: Production | IE | ME", 1.0, 3.8, 11, 0.4, size=16, color=WHITE)
T(sl, "Tujuan: Membangun satu layar pemantauan otomatis untuk semua kebutuhan data produksi.",
  1.0, 4.4, 10, 1.0, size=12, color=CYAN)


# ==========================================
# 2. MASALAH & SOLUSI
# ==========================================
sl = new_slide()
title_bar(sl, "Masalah & Solusi", "Mengapa kita butuh Dashboard Material Ready?")
# Before
R(sl, 0.5, 1.3, 5.8, 5.5, fill=RGBColor(0xFF,0xEB,0xEE), line=RED, lw=2, radius=True)
R(sl, 0.5, 1.3, 5.8, 0.6, fill=RED, radius=True)
T(sl, "KONDISI SAAT INI (MANUAL)", 0.5, 1.45, 5.8, 0.4, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
before = [
    "Operator harus buka APS, Engage, dan CIUROX secara terpisah setiap hari.",
    "Data digabung secara manual di Excel (rawan salah & lama).",
    "Supervisor baru tahu kondisi produksi setelah laporan dikirim (tidak real-time)."
]
bullet_list(sl, before, 0.8, 2.3, 5.2, size=13, gap=1.0, color=DARK, dot_color=RED)

# After
R(sl, 7.0, 1.3, 5.8, 5.5, fill=RGBColor(0xE8,0xF5,0xE9), line=GREEN, lw=2, radius=True)
R(sl, 7.0, 1.3, 5.8, 0.6, fill=GREEN, radius=True)
T(sl, "SOLUSI DASHBOARD (OTOMATIS)", 7.0, 1.45, 5.8, 0.4, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
after = [
    "Sistem otomatis menarik data dari APS, Engage & CIUROX setiap jam.",
    "Data langsung tampil di 1 layar web, siap diakses semua orang.",
    "Supervisor bisa memantau kondisi lapangan secara Real-Time."
]
bullet_list(sl, after, 7.3, 2.3, 5.2, size=13, gap=1.0, color=DARK, dot_color=GREEN)

T(sl, "→", 6.4, 3.5, 0.5, 0.8, size=40, bold=True, color=CYAN)
footer(sl, 2)


# ==========================================
# 3. SUMBER DATA & CARA KERJA
# ==========================================
sl = new_slide()
title_bar(sl, "Darimana Data Berasal?", "Sistem menggabungkan 3 sumber utama secara otomatis")

NAVY_LIGHT = RGBColor(0x2E,0x55,0x8A)
R(sl, 1.0, 1.5, 3.2, 1.5, fill=NAVY_LIGHT, radius=True)
T(sl, "APS", 1.0, 1.7, 3.2, 0.4, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
T(sl, "Master Data Order\nTarget Produksi\nJadwal Pengiriman", 1.0, 2.2, 3.2, 0.8, size=11, color=WHITE, align=PP_ALIGN.CENTER)

R(sl, 5.0, 1.5, 3.2, 1.5, fill=TEAL, radius=True)
T(sl, "Engage", 5.0, 1.7, 3.2, 0.4, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
T(sl, "Stok Gudang (In/Out)\nOutput Aktual", 5.0, 2.3, 3.2, 0.8, size=11, color=WHITE, align=PP_ALIGN.CENTER)

R(sl, 9.0, 1.5, 3.2, 1.5, fill=RGBColor(0x1B,0x5E,0x20), radius=True)
T(sl, "CIUROX", 9.0, 1.7, 3.2, 0.4, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
T(sl, "Status Kesiapan\nAksesoris (Completed)", 9.0, 2.3, 3.2, 0.8, size=11, color=WHITE, align=PP_ALIGN.CENTER)

T(sl, "↓", 6.4, 3.2, 0.5, 0.5, size=30, bold=True, color=CYAN)

R(sl, 3.0, 3.8, 7.2, 1.8, fill=GRAY, line=NAVY, lw=2, radius=True)
T(sl, "DASHBOARD MATERIAL READY", 3.0, 4.2, 7.2, 0.5, size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
T(sl, "Sistem melakukan pengolahan & perhitungan secara otomatis di background.", 3.0, 4.8, 7.2, 0.5, size=12, color=DARK, align=PP_ALIGN.CENTER)

footer(sl, 3)


# ==========================================
# 4. KONSEP TAMPILAN & CARA PERHITUNGAN
# ==========================================
sl = new_slide()
title_bar(sl, "Konsep Tampilan & Rumus Perhitungan", "Informasi apa yang ditampilkan dan dari mana angkanya?")

# Konsep Tampilan (Kiri)
R(sl, 0.4, 1.2, 5.5, 5.8, fill=RGBColor(0xEC,0xEF,0xF1), line=NAVY, lw=1)
R(sl, 0.4, 1.2, 5.5, 0.4, fill=NAVY)
T(sl, "Mockup Dashboard", 0.5, 1.28, 5.0, 0.3, size=11, bold=True, color=WHITE)
# mock kpi
for i, (lbl, val, c) in enumerate([("Total Output", "2.840", NAVY_LIGHT), ("Balance Qty", "460", AMBER)]):
    R(sl, 0.6 + i*2.6, 1.8, 2.4, 0.9, fill=c, radius=True)
    T(sl, val, 0.6 + i*2.6, 1.9, 2.4, 0.4, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    T(sl, lbl, 0.6 + i*2.6, 2.35, 2.4, 0.3, size=9, color=WHITE, align=PP_ALIGN.CENTER)
# mock table
R(sl, 0.6, 2.9, 5.1, 1.8, fill=WHITE, line=GRAY_MID, lw=1)
T(sl, "Material Ready (Top Priority)", 0.7, 3.0, 4.0, 0.3, size=9, bold=True, color=NAVY)
T(sl, "Order A | END AUG | READY\nOrder B | END AUG | READY\nOrder C | MID SEP | PROSES", 0.7, 3.4, 4.0, 1.0, size=9)
# mock chart
R(sl, 0.6, 4.9, 5.1, 1.8, fill=WHITE, line=GRAY_MID, lw=1)
T(sl, "Demand vs Output", 0.7, 5.0, 4.0, 0.3, size=9, bold=True, color=NAVY)
R(sl, 1.0, 5.5, 4.0, 0.8, fill=RGBColor(0xE3,0xF2,0xFD))

# Rumus Perhitungan (Kanan)
T(sl, "Cara Sistem Menghitung (Logic)", 6.3, 1.3, 6.0, 0.4, size=14, bold=True, color=NAVY)

calcs = [
    ("Total Output", "Dari Engage (Transaksi barang keluar/outflow di gudang)."),
    ("Balance (Sisa Target)", "Target Produksi (APS) dikurangi Total Output (Engage)."),
    ("Demand Harian", "Balance dibagi Sisa Hari Kerja.\n(Memberi tahu berapa minimal yg harus diproduksi per hari)."),
    ("Material Ready (RTL)", "Order dinyatakan READY jika memenuhi 3 Syarat:\n"
                             "1. Terjadwal di APS\n"
                             "2. Barang Inflow sudah ada di Engage\n"
                             "3. Status aksesoris COMPLETED di CIUROX"),
    ("Prioritas Order", "Diurutkan otomatis dari tanggal pengiriman terdekat di APS (Misal: MID AUG lebih dulu dari END AUG).")
]

y = 1.9
for title, desc in calcs:
    R(sl, 6.3, y, 6.6, 0.9, fill=GRAY, line=RGBColor(0xDD,0xDD,0xDD), lw=1, radius=True)
    T(sl, title, 6.4, y+0.1, 6.4, 0.3, size=11, bold=True, color=TEAL)
    T(sl, desc, 6.4, y+0.4, 6.4, 0.4, size=9.5, color=DARK)
    y += 1.05

footer(sl, 4)


# ==========================================
# 5. DISKUSI & REVIEW KEBUTUHAN
# ==========================================
sl = new_slide()
title_bar(sl, "Diskusi & Review Kebutuhan", "Sampaikan kebutuhan harian tim Anda")

PROD_CLR = RGBColor(0x1B,0x5E,0x20)
IE_CLR = RGBColor(0x01,0x57,0x9B)
ME_CLR = RGBColor(0x4A,0x14,0x8C)
cols = [
    ("Production", PROD_CLR, "Apakah urutan prioritas sudah cukup membantu? Ada info khusus yg sering ditanyakan ke gudang?"),
    ("Industrial Engineering", IE_CLR, "Apakah format Demand vs Output sudah sesuai standar IE untuk mengukur efisiensi?"),
    ("Manufacturing Engineering", ME_CLR, "Apakah info kesiapan aksesoris (CIUROX) yang digabung dengan Engage sudah cukup untuk planning loading?")
]

for i, (dept, clr, question) in enumerate(cols):
    l = 0.5 + i * 4.2
    R(sl, l, 1.5, 3.8, 3.0, fill=WHITE, line=clr, lw=2, radius=True)
    R(sl, l, 1.5, 3.8, 0.6, fill=clr, radius=True)
    T(sl, dept, l+0.1, 1.6, 3.6, 0.4, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    T(sl, question, l+0.2, 2.3, 3.4, 1.8, size=12, color=DARK)

# Next steps
R(sl, 0.5, 5.2, 12.2, 1.6, fill=RGBColor(0xE3,0xF2,0xFD), line=NAVY_LIGHT, lw=2, radius=True)
T(sl, "Langkah Selanjutnya (Next Steps)", 0.7, 5.35, 11.8, 0.4, size=13, bold=True, color=NAVY)
steps = [
    "Kumpulkan masukan akhir dari setiap departemen hari ini.",
    "Tim IT akan merampungkan Prototype Dashboard.",
    "UAT (Uji Coba) bersama sebelum Go-Live resmi digunakan."
]
bullet_list(sl, steps, 0.9, 5.8, 11.0, size=12, gap=0.35, color=DARK, dot_color=CYAN)

footer(sl, 5)


# ==========================================
# 6. PENUTUP
# ==========================================
sl = new_slide()
R(sl, 0, 0, 13.33, 7.5, fill=NAVY)
T(sl, "Terima Kasih!", 1.0, 2.5, 11, 1.0, size=52, bold=True, color=WHITE)
T(sl, "Dengan dashboard ini, kita akan bekerja dengan SATU DATA YANG SAMA, LEBIH CEPAT, dan LEBIH AKURAT.",
  1.0, 3.7, 11, 1.0, size=16, color=CYAN)
footer(sl, 6)

# SAVE
out_path = r"c:\xampp\htdocs\dashboard_gm\IRL3_Presentasi_Konsep_Dashboard.pptx"
prs.save(out_path)
print(f"[OK] Saved: {out_path} ({TOTAL} slides)")
