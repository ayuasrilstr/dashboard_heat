# coding: utf-8
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

NAVY_HEADER = RGBColor(0x13, 0x52, 0x8A)
NAVY_PILL   = RGBColor(0x0E, 0x4D, 0x82)
CYAN_PILL   = RGBColor(0x02, 0x88, 0xD1)
GREEN_PILL  = RGBColor(0x43, 0xA0, 0x47)
AMBER_PILL  = RGBColor(0xFB, 0x8C, 0x00)
LIGHT_BLUE  = RGBColor(0xE1, 0xF5, 0xFE)
LIGHT_GREEN = RGBColor(0xE8, 0xF5, 0xE9)
LIGHT_AMBER = RGBColor(0xFF, 0xF8, 0xE1)
BORDER_BOX  = RGBColor(0x90, 0xA4, 0xAE)
DARK_TEXT   = RGBColor(0x21, 0x25, 0x29)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
sl = prs.slides.add_slide(prs.slide_layouts[6])

def R(l, t, w, h, fill=WHITE, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    return s

def T(text, l, t, w, h, size=10, bold=False, color=DARK_TEXT, align=PP_ALIGN.LEFT, italic=False):
    tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb

# Outer frame
R(0.1, 0.1, 13.133, 7.3, fill=WHITE, line=RGBColor(0xCF, 0xD8, 0xDC), lw=1)

# Top Title Banner
R(0.2, 0.2, 10.4, 0.52, fill=NAVY_HEADER)
T('IRL 5 : DEVELOPMENT DASHBOARD, INTEGRASI DATABASE, SERTA PENGUJIAN KONEKSI DATA REAL-TIME',
  0.35, 0.28, 10.1, 0.38, size=13, bold=True, color=WHITE)

# ===============================================================
# SECTION 1 (LEFT): INTEGRASI DATABASE
# ===============================================================
R(2.0, 0.95, 2.5, 0.36, fill=NAVY_PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('INTEGRASI DATABASE', 2.0, 1.0, 2.5, 0.28, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# 3 Ingestion Source Boxes
R(0.4, 1.5, 1.65, 0.45, fill=WHITE, line=BORDER_BOX, lw=1)
T('Accessories RPA\n(CIUROX Web)', 0.4, 1.52, 1.65, 0.4, size=8, bold=True, align=PP_ALIGN.CENTER)

R(0.4, 2.05, 1.65, 0.4, fill=WHITE, line=BORDER_BOX, lw=0.8)
T('Auto CONTROLIST.xlsx\n(Playwright Runner)', 0.4, 2.07, 1.65, 0.35, size=7.5, align=PP_ALIGN.CENTER)

R(2.35, 1.5, 1.65, 0.45, fill=WHITE, line=BORDER_BOX, lw=1)
T('APS RPA\n(IOS-APS UNC Server)', 2.35, 1.52, 1.65, 0.4, size=8, bold=True, align=PP_ALIGN.CENTER)

R(2.35, 2.05, 1.65, 0.4, fill=WHITE, line=BORDER_BOX, lw=0.8)
T('Export JO.xlsx\n(Win32 GUI Automation)', 2.35, 2.07, 1.65, 0.35, size=7.5, align=PP_ALIGN.CENTER)

R(4.3, 1.5, 1.65, 0.45, fill=WHITE, line=BORDER_BOX, lw=1)
T('Engage Warehouse\n(Storage 32 & 32a)', 4.3, 1.52, 1.65, 0.4, size=8, bold=True, align=PP_ALIGN.CENTER)

R(4.3, 2.05, 1.65, 0.4, fill=WHITE, line=BORDER_BOX, lw=0.8)
T('sync_engage_transactions.php\n(PHP CLI Sync Daemon)', 4.3, 2.07, 1.65, 0.35, size=7, align=PP_ALIGN.CENTER)

# Arrow from sources to ETL Engine
s_arrow = sl.shapes.add_connector(1, Inches(3.175), Inches(2.5), Inches(3.175), Inches(2.78))
s_arrow.line.color.rgb = NAVY_PILL
s_arrow.line.width = Pt(1.5)

# Central Processing Box: Data Pipeline & ETL Engine
R(1.4, 2.8, 3.55, 0.65, fill=WHITE, line=NAVY_PILL, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('Pipeline ETL & Data Cleansing Engine', 1.4, 2.83, 3.55, 0.25, size=9.5, bold=True, color=NAVY_PILL, align=PP_ALIGN.CENTER)
T('Streaming Native XML Parser (No-Leak) • Order Normalizer • Sanitasi Null/Date',
  1.45, 3.12, 3.45, 0.28, size=7.5, color=DARK_TEXT, align=PP_ALIGN.CENTER)

# Arrow to Database
s_arrow2 = sl.shapes.add_connector(1, Inches(3.175), Inches(3.5), Inches(3.175), Inches(3.78))
s_arrow2.line.color.rgb = NAVY_PILL
s_arrow2.line.width = Pt(1.5)

# Database Cylinder / Box
R(2.0, 3.8, 2.35, 0.55, fill=WHITE, line=NAVY_PILL, lw=1.5, shape=MSO_SHAPE.CAN)
T('Database MySQL\n(db_dashboardgm) & JSON Cache', 2.0, 3.92, 2.35, 0.4, size=8.5, bold=True, color=NAVY_PILL, align=PP_ALIGN.CENTER)

# Dashed labels
T('- - - > tb_engage_transactions (Staging Mutasi Live)', 4.45, 3.75, 2.2, 0.25, size=7, color=DARK_TEXT)
T('- - - > tb_engage_archieve (Auto-Rotasi 90 Hari)', 4.45, 3.98, 2.2, 0.25, size=7, color=DARK_TEXT)
T('- - - > engage_daily_history & dashboard_heat_history', 4.45, 4.20, 2.2, 0.25, size=7, color=DARK_TEXT)

# Keterangan
T('Keterangan :', 0.4, 4.65, 5.8, 0.25, size=9.5, bold=True, color=NAVY_PILL)

ket_items = [
    ('Dual Data Source Handshake:', ' Mengintegrasikan Active Record MySQL dengan native XML stream reader untuk Excel (proses ribuan baris < 1 detik tanpa memory leak).'),
    ('Atomic Upsert & Deduplikasi:', ' Penyeragaman transaksi Engage dengan pencegahan duplikasi data melalui komposit unique key (udef_3 / udef_4).'),
    ('Automated Retention & Archiving:', ' Pemisahan data mutasi harian dan arsip lampau > 90 hari untuk menjaga latency query MySQL selalu di bawah 15ms.'),
    ('Low-Latency Local JSON Cache:', ' Konfigurasi dinamis (SMV per style, kalender kerja/shift, analytics) dicache lokal untuk meminimalkan beban I/O database.')
]

ky = 4.95
for title, desc in ket_items:
    tb = sl.shapes.add_textbox(Inches(0.4), Inches(ky), Inches(5.8), Inches(0.48))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r0 = p.add_run()
    r0.text = '• ' + title
    r0.font.bold = True
    r0.font.size = Pt(8)
    r0.font.color.rgb = DARK_TEXT
    r1 = p.add_run()
    r1.text = desc
    r1.font.bold = False
    r1.font.size = Pt(7.8)
    r1.font.color.rgb = DARK_TEXT
    ky += 0.52

# ===============================================================
# SECTION 2 (TOP RIGHT): PENGUJIAN KONEKSI DATA REAL-TIME
# ===============================================================
R(8.8, 0.95, 3.2, 0.36, fill=NAVY_PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('PENGUJIAN KONEKSI DATA REAL-TIME', 8.8, 1.0, 3.2, 0.28, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# 3 Column Sub-Header Pills
R(6.5, 1.45, 1.9, 0.32, fill=CYAN_PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('KOMPONEN KONEKSI', 6.5, 1.5, 1.9, 0.22, size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

R(8.6, 1.45, 2.0, 0.32, fill=GREEN_PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('METODE PENGUJIAN', 8.6, 1.5, 2.0, 0.22, size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

R(10.8, 1.45, 2.1, 0.32, fill=AMBER_PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('HASIL & STATUS VALIDASI', 10.8, 1.5, 2.1, 0.22, size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# Row 1: MySQL Database Connection
R(6.5, 1.88, 1.9, 0.52, fill=LIGHT_BLUE, line=BORDER_BOX, lw=0.8)
T('Database MySQL\n(db_dashboardgm)', 6.55, 1.95, 1.8, 0.4, size=8, bold=True, align=PP_ALIGN.CENTER)

arr1 = sl.shapes.add_connector(1, Inches(8.42), Inches(2.14), Inches(8.58), Inches(2.14))
arr1.line.color.rgb = NAVY_PILL; arr1.line.width = Pt(1.5)

R(8.6, 1.88, 2.0, 0.52, fill=LIGHT_GREEN, line=BORDER_BOX, lw=0.8)
T('Stress test query Active Record & pooling (100 concurrent req)', 8.65, 1.92, 1.9, 0.45, size=7.5, align=PP_ALIGN.CENTER)

arr2 = sl.shapes.add_connector(1, Inches(10.62), Inches(2.14), Inches(10.78), Inches(2.14))
arr2.line.color.rgb = NAVY_PILL; arr2.line.width = Pt(1.5)

R(10.8, 1.88, 2.1, 0.52, fill=LIGHT_AMBER, line=BORDER_BOX, lw=0.8)
T('Latency < 15ms, Zero Data Loss, Integritas Relasi 100% [PASS]', 10.85, 1.92, 2.0, 0.45, size=7.5, bold=True, align=PP_ALIGN.CENTER)

# Row 2: Pipeline RPA Master Scheduler
R(6.5, 2.50, 1.9, 0.52, fill=LIGHT_BLUE, line=BORDER_BOX, lw=0.8)
T('Pipeline Otomasi RPA\n(RPA_Master.exe)', 6.55, 2.57, 1.8, 0.4, size=8, bold=True, align=PP_ALIGN.CENTER)

arr3 = sl.shapes.add_connector(1, Inches(8.42), Inches(2.76), Inches(8.58), Inches(2.76))
arr3.line.color.rgb = NAVY_PILL; arr3.line.width = Pt(1.5)

R(8.6, 2.50, 2.0, 0.52, fill=LIGHT_GREEN, line=BORDER_BOX, lw=0.8)
T('Simulasi cron 90 menit & on-demand API (/api/run-download)', 8.65, 2.54, 1.9, 0.45, size=7.5, align=PP_ALIGN.CENTER)

arr4 = sl.shapes.add_connector(1, Inches(10.62), Inches(2.76), Inches(10.78), Inches(2.76))
arr4.line.color.rgb = NAVY_PILL; arr4.line.width = Pt(1.5)

R(10.8, 2.50, 2.1, 0.52, fill=LIGHT_AMBER, line=BORDER_BOX, lw=0.8)
T('Process-Lock aktif, Auto-Recovery, Zero Race Condition [PASS]', 10.85, 2.54, 2.0, 0.45, size=7.5, bold=True, align=PP_ALIGN.CENTER)

# Row 3: Client Real-Time Sync
R(6.5, 3.12, 1.9, 0.52, fill=LIGHT_BLUE, line=BORDER_BOX, lw=0.8)
T('Sinkronisasi Client Real-Time', 6.55, 3.19, 1.8, 0.4, size=8, bold=True, align=PP_ALIGN.CENTER)

arr5 = sl.shapes.add_connector(1, Inches(8.42), Inches(3.38), Inches(8.58), Inches(3.38))
arr5.line.color.rgb = NAVY_PILL; arr5.line.width = Pt(1.5)

R(8.6, 3.12, 2.0, 0.52, fill=LIGHT_GREEN, line=BORDER_BOX, lw=0.8)
T('Uji wall-clock refresh (:00 & :30) + Pre-Hydration payload', 8.65, 3.16, 1.9, 0.45, size=7.5, align=PP_ALIGN.CENTER)

arr6 = sl.shapes.add_connector(1, Inches(10.62), Inches(3.38), Inches(10.78), Inches(3.38))
arr6.line.color.rgb = NAVY_PILL; arr6.line.width = Pt(1.5)

R(10.8, 3.12, 2.1, 0.52, fill=LIGHT_AMBER, line=BORDER_BOX, lw=0.8)
T('Render instan 0ms delay, payload ~35KB, Uptime 24/7 TV [PASS]', 10.85, 3.16, 2.0, 0.45, size=7.5, bold=True, align=PP_ALIGN.CENTER)

# ===============================================================
# SECTION 3 (BOTTOM RIGHT): DEVELOPMENT DASHBOARD
# ===============================================================
R(8.8, 3.85, 3.2, 0.36, fill=NAVY_PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
T('DEVELOPMENT DASHBOARD', 8.8, 3.9, 3.2, 0.28, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

dev_bullets = [
    'Arsitektur SPA Berbasis CodeIgniter 3: Frontend terpadu Vanilla JS + CSS3 yang ringan, responsif, dan optimal untuk layar monitor TV & browser desktop.',
    'Instant Server-Side Pre-Hydration: Injeksi data via INITIAL_DASHBOARD_PAYLOAD saat HTML disajikan, mengeliminasi flickering / delay loading.',
    'Multi-Delivery Scope Dynamic Switcher: Fleksibilitas memilih rentang 1, 2, 4, atau 6 delivery dengan persistensi otomatis berbasis browser cookie.',
    'Panel Admin Terpadu GM (Protected Session): Manajemen terproteksi untuk konfigurasi Style SMV, Toggle Direct Aktual, Kalender Hari Kerja, dan visibilitas metrik.',
    'Mesin Evaluasi Analytics & CAP: Deteksi otomatis status risiko operasional (Good / Watch / Risk) serta formulasi Corrective Action Plan harian.',
    'Dynamic Excel Report Exporter: Fitur ekspor data tabel Material To Load ke format spreadsheet secara dinamis menyesuaikan filter periode delivery aktif.',
    'Wall-Clock Aligned Auto-Refresh: Pembaruan data otomatis tepat pada menit :00 dan :30 setiap jam secara sinkron dengan jam dinding pabrik & status live pulse.',
    'Sistem Diagnostik & Centralized Logging: Pemantauan kesehatan koneksi data terintegrasi dengan file log terpusat (scheduler.log) dan opsi trigger manual.'
]

dy = 4.30
for b in dev_bullets:
    tb = sl.shapes.add_textbox(Inches(6.5), Inches(dy), Inches(6.5), Inches(0.36))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = '• ' + b
    r.font.size = Pt(8)
    r.font.color.rgb = DARK_TEXT
    dy += 0.37

pptx_path = 'IRL5_Development_dan_Pengujian_1slide.pptx'
prs.save(pptx_path)
print(f'[OK] Successfully saved {pptx_path}')
