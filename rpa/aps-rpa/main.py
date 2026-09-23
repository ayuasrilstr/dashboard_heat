import argparse
import ctypes
import json
import logging
import os
import subprocess
import shutil
import sys
import time
from ctypes import wintypes
from datetime import datetime, timedelta
from pathlib import Path

try:
    from PIL import ImageChops, ImageGrab, ImageStat
except ImportError:
    ImageGrab = None
    ImageChops = None
    ImageStat = None


ROOT = Path(__file__).resolve().parent
DEFAULT_CONFIG = ROOT / "config.json"
DEFAULT_ENV = ROOT / ".env"

user32 = ctypes.WinDLL("user32", use_last_error=True)

EnumWindowsProc = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)

user32.EnumWindows.argtypes = [EnumWindowsProc, wintypes.LPARAM]
user32.EnumWindows.restype = wintypes.BOOL
user32.GetWindowTextLengthW.argtypes = [wintypes.HWND]
user32.GetWindowTextLengthW.restype = ctypes.c_int
user32.GetWindowTextW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
user32.GetWindowTextW.restype = ctypes.c_int
user32.IsWindowVisible.argtypes = [wintypes.HWND]
user32.IsWindowVisible.restype = wintypes.BOOL
user32.SetForegroundWindow.argtypes = [wintypes.HWND]
user32.SetForegroundWindow.restype = wintypes.BOOL
user32.ShowWindow.argtypes = [wintypes.HWND, ctypes.c_int]
user32.ShowWindow.restype = wintypes.BOOL
user32.keybd_event.argtypes = [wintypes.BYTE, wintypes.BYTE, wintypes.DWORD, ctypes.POINTER(ctypes.c_ulong)]
user32.mouse_event.argtypes = [wintypes.DWORD, wintypes.DWORD, wintypes.DWORD, wintypes.DWORD, ctypes.POINTER(ctypes.c_ulong)]
user32.SetCursorPos.argtypes = [ctypes.c_int, ctypes.c_int]
user32.SetCursorPos.restype = wintypes.BOOL
user32.GetSystemMetrics.argtypes = [ctypes.c_int]
user32.GetSystemMetrics.restype = ctypes.c_int
user32.GetForegroundWindow.restype = wintypes.HWND
user32.OpenClipboard.argtypes = [wintypes.HWND]
user32.OpenClipboard.restype = wintypes.BOOL
user32.EmptyClipboard.restype = wintypes.BOOL
user32.SetClipboardData.argtypes = [wintypes.UINT, wintypes.HANDLE]
user32.SetClipboardData.restype = wintypes.HANDLE
user32.CloseClipboard.restype = wintypes.BOOL
user32.VkKeyScanW.argtypes = [ctypes.c_wchar]
user32.VkKeyScanW.restype = ctypes.c_short
user32.MapVirtualKeyW.argtypes = [wintypes.UINT, wintypes.UINT]
user32.MapVirtualKeyW.restype = wintypes.UINT

user32.GetWindowThreadProcessId.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.DWORD)]
user32.GetWindowThreadProcessId.restype = wintypes.DWORD
user32.AttachThreadInput.argtypes = [wintypes.DWORD, wintypes.DWORD, wintypes.BOOL]
user32.AttachThreadInput.restype = wintypes.BOOL
user32.BringWindowToTop.argtypes = [wintypes.HWND]
user32.BringWindowToTop.restype = wintypes.BOOL
user32.SetWindowPos.argtypes = [wintypes.HWND, wintypes.HWND, ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int, wintypes.UINT]
user32.SetWindowPos.restype = wintypes.BOOL
user32.PostMessageW.argtypes = [wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
user32.PostMessageW.restype = wintypes.BOOL

kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
kernel32.GetCurrentThreadId.restype = wintypes.DWORD
kernel32.GlobalAlloc.argtypes = [wintypes.UINT, ctypes.c_size_t]
kernel32.GlobalAlloc.restype = wintypes.HGLOBAL
kernel32.GlobalLock.argtypes = [wintypes.HGLOBAL]
kernel32.GlobalLock.restype = wintypes.LPVOID
kernel32.GlobalUnlock.argtypes = [wintypes.HGLOBAL]
kernel32.GlobalUnlock.restype = wintypes.BOOL

KEYEVENTF_KEYUP = 0x0002
MOUSEEVENTF_LEFTDOWN = 0x0002
MOUSEEVENTF_LEFTUP = 0x0004
SM_CXSCREEN = 0
SM_CYSCREEN = 1
GMEM_MOVEABLE = 0x0002
CF_UNICODETEXT = 13
SW_SHOWNORMAL = 1
SW_MAXIMIZE = 3
SW_RESTORE = 9
HWND_TOPMOST = -1
HWND_NOTOPMOST = -2
SWP_NOSIZE = 0x0001
SWP_NOMOVE = 0x0002
SWP_SHOWWINDOW = 0x0040
WM_CLOSE = 0x0010


class RECT(ctypes.Structure):
    _fields_ = [
        ("left", wintypes.LONG),
        ("top", wintypes.LONG),
        ("right", wintypes.LONG),
        ("bottom", wintypes.LONG),
    ]


wintypes.RECT = RECT
user32.GetWindowRect.argtypes = [wintypes.HWND, ctypes.POINTER(RECT)]
user32.GetWindowRect.restype = wintypes.BOOL

VK = {
    "alt": 0x12,
    "backspace": 0x08,
    "ctrl": 0x11,
    "delete": 0x2E,
    "down": 0x28,
    "end": 0x23,
    "enter": 0x0D,
    "esc": 0x1B,
    "f1": 0x70,
    "f2": 0x71,
    "f3": 0x72,
    "f4": 0x73,
    "f5": 0x74,
    "f6": 0x75,
    "f7": 0x76,
    "f8": 0x77,
    "f9": 0x78,
    "f10": 0x79,
    "f11": 0x7A,
    "f12": 0x7B,
    "home": 0x24,
    "left": 0x25,
    "pagedown": 0x22,
    "pageup": 0x21,
    "right": 0x27,
    "shift": 0x10,
    "space": 0x20,
    "tab": 0x09,
    "up": 0x26,
}

for digit in "0123456789":
    VK[digit] = ord(digit)
for letter in "abcdefghijklmnopqrstuvwxyz":
    VK[letter] = ord(letter.upper())


def load_config(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(f"Config tidak ditemukan: {path}")
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" not in line:
            raise ValueError(f"Format .env tidak valid di baris {line_number}: {raw_line}")
        name, value = line.split("=", 1)
        name = name.strip()
        value = value.strip().strip('"').strip("'")
        if not name:
            raise ValueError(f"Nama environment kosong di .env baris {line_number}")
        os.environ.setdefault(name, value)


def setup_logging() -> Path:
    log_dir = ROOT / "logs"
    log_dir.mkdir(exist_ok=True)
    log_path = log_dir / f"run-{datetime.now():%Y%m%d-%H%M%S}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(log_path, encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )
    return log_path


def window_title(hwnd: int) -> str:
    length = user32.GetWindowTextLengthW(hwnd)
    if length == 0:
        return ""
    buffer = ctypes.create_unicode_buffer(length + 1)
    user32.GetWindowTextW(hwnd, buffer, length + 1)
    return buffer.value


def visible_windows() -> list[tuple[int, str]]:
    result: list[tuple[int, str]] = []

    @EnumWindowsProc
    def callback(hwnd, _):
        if user32.IsWindowVisible(hwnd):
            title = window_title(hwnd).strip()
            if title:
                result.append((int(hwnd), title))
        return True

    user32.EnumWindows(callback, 0)
    return result


def title_matches(title: str, title_contains: str | list[str] | None) -> bool:
    if not title_contains:
        return True
    options = title_contains if isinstance(title_contains, list) else [title_contains]
    title_lower = title.lower()

    excel_lookups = {"excel", "microsoft excel", "worksheet", "book", "wps", "wps office", "spreadsheets"}
    if any(str(opt).lower() in excel_lookups for opt in options):
        if any(ign in title_lower for ign in ["photo", "photos", "image", "viewer", "picture"]):
            return False

    return any(str(option).lower() in title_lower for option in options)


def find_window(title_contains: str | list[str] | None) -> tuple[int, str] | None:
    windows = visible_windows()
    for hwnd, title in windows:
        if title_matches(title, title_contains):
            return hwnd, title
    return None


def maximize_window(hwnd: int) -> None:
    if hwnd:
        user32.ShowWindow(hwnd, SW_MAXIMIZE)
        time.sleep(0.4)


def focus_window(hwnd: int, maximize: bool = False) -> None:
    if not hwnd:
        return
    fg_hwnd = user32.GetForegroundWindow()
    if fg_hwnd == hwnd and not maximize:
        return
    try:
        user32.keybd_event(0x12, 0, 0, 0)
        user32.keybd_event(0x12, 0, 2, 0)
    except Exception:
        pass
    if maximize:
        maximize_window(hwnd)
    else:
        if user32.IsIconic(hwnd):
            user32.ShowWindow(hwnd, SW_RESTORE)
        elif user32.GetForegroundWindow() != hwnd:
            user32.ShowWindow(hwnd, SW_SHOWNORMAL)
    user32.BringWindowToTop(hwnd)
    user32.SetForegroundWindow(hwnd)

    fg_hwnd = user32.GetForegroundWindow()
    if fg_hwnd != hwnd:
        fg_thread = user32.GetWindowThreadProcessId(fg_hwnd, None)
        cur_thread = kernel32.GetCurrentThreadId()
        if fg_thread and fg_thread != cur_thread:
            user32.AttachThreadInput(cur_thread, fg_thread, True)
            user32.BringWindowToTop(hwnd)
            user32.SetForegroundWindow(hwnd)
            user32.AttachThreadInput(cur_thread, fg_thread, False)

    try:
        user32.SetWindowPos(hwnd, HWND_TOPMOST, 0, 0, 0, 0, SWP_NOMOVE | SWP_NOSIZE | SWP_SHOWWINDOW)
        user32.SetWindowPos(hwnd, HWND_NOTOPMOST, 0, 0, 0, 0, SWP_NOMOVE | SWP_NOSIZE | SWP_SHOWWINDOW)
    except Exception:
        pass

    time.sleep(0.3)


def press_key(name: str) -> None:
    key = VK.get(name.lower())
    if key is None:
        raise ValueError(f"Tombol tidak dikenal: {name}")
    scan = user32.MapVirtualKeyW(key, 0)
    user32.keybd_event(key, scan, 0, None)
    time.sleep(0.03)
    user32.keybd_event(key, scan, KEYEVENTF_KEYUP, None)


def press_hotkey(keys: list[str]) -> None:
    codes = []
    for key_name in keys:
        key = VK.get(key_name.lower())
        if key is None:
            raise ValueError(f"Tombol tidak dikenal: {key_name}")
        codes.append(key)
    for key in codes:
        scan = user32.MapVirtualKeyW(key, 0)
        user32.keybd_event(key, scan, 0, None)
        time.sleep(0.03)
    for key in reversed(codes):
        scan = user32.MapVirtualKeyW(key, 0)
        user32.keybd_event(key, scan, KEYEVENTF_KEYUP, None)
        time.sleep(0.03)


def press_configured_keys(keys: list[str]) -> None:
    if len(keys) == 1:
        press_key(str(keys[0]))
    else:
        press_hotkey([str(key) for key in keys])


def find_popup_window(title_contains: str | list[str] | None, ignored_hwnd: int | None = None) -> tuple[int, str] | None:
    for hwnd, title in visible_windows():
        if title_matches(title, title_contains):
            return hwnd, title
    return None


def handle_optional_popup(popup_config: dict, ignored_hwnd: int | None = None) -> bool:
    title_contains = popup_config.get("title_contains", ["Prompt", "Message", "Warning", "Error", "Confirm"])
    found = find_popup_window(title_contains, ignored_hwnd)
    if not found:
        return False

    hwnd, title = found
    logging.info("Popup terdeteksi: %s", title)
    focus_window(hwnd)
    press_configured_keys(popup_config.get("keys", ["enter"]))
    time.sleep(float(popup_config.get("after_seconds", 1)))
    return True


def wait_for_window(title_contains: str | list[str] | None, timeout_seconds: int, popup_config: dict | None = None) -> tuple[int, str]:
    deadline = time.time() + timeout_seconds
    last_popup_check = 0.0
    while time.time() < deadline:
        found = find_window(title_contains)
        if found:
            return found
        if popup_config and (time.time() - last_popup_check >= float(popup_config.get("check_interval_seconds", 2))):
            last_popup_check = time.time()
            handle_optional_popup(popup_config)
        time.sleep(0.5)
    visible = ", ".join(title for _, title in visible_windows()[:10])
    raise TimeoutError(f"Window tidak ditemukan. Window terlihat: {visible}")


def wait_window_disappear(title_contains: str | list[str], timeout_seconds: int = 30) -> bool:
    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        if not find_window(title_contains):
            return True
        time.sleep(0.5)
    logging.warning("Window %s belum hilang setelah %s detik.", title_contains, timeout_seconds)
    return False


def wait_for_any_window(title_options: list[str], timeout_seconds: int) -> tuple[int, str] | None:
    deadline = time.time() + timeout_seconds
    lowered = [title.lower() for title in title_options]
    while time.time() < deadline:
        for hwnd, title in visible_windows():
            title_lower = title.lower()
            if any(option in title_lower for option in lowered):
                return hwnd, title
        time.sleep(0.5)
    return None


def click(x: int, y: int) -> None:
    screen_width = user32.GetSystemMetrics(SM_CXSCREEN)
    screen_height = user32.GetSystemMetrics(SM_CYSCREEN)
    if not (0 <= x < screen_width and 0 <= y < screen_height):
        raise ValueError(f"Koordinat klik di luar layar: x={x}, y={y}, layar={screen_width}x{screen_height}")
    user32.SetCursorPos(x, y)
    time.sleep(0.1)
    user32.mouse_event(MOUSEEVENTF_LEFTDOWN, 0, 0, 0, None)
    time.sleep(0.08)
    user32.mouse_event(MOUSEEVENTF_LEFTUP, 0, 0, 0, None)


def type_text(text: str) -> None:
    for char in text:
        if char == "\n":
            press_key("enter")
        elif char == "\t":
            press_key("tab")
        else:
            vk_scan = user32.VkKeyScanW(char)
            if vk_scan == -1:
                raise ValueError(f"Karakter tidak bisa diketik otomatis: {char!r}")
            vk = vk_scan & 0xFF
            scan = user32.MapVirtualKeyW(vk, 0)
            shift_state = (vk_scan >> 8) & 0xFF
            if shift_state & 1:
                shift_scan = user32.MapVirtualKeyW(VK["shift"], 0)
                user32.keybd_event(VK["shift"], shift_scan, 0, None)
            user32.keybd_event(vk, scan, 0, None)
            time.sleep(0.04)
            user32.keybd_event(vk, scan, KEYEVENTF_KEYUP, None)
            if shift_state & 1:
                user32.keybd_event(VK["shift"], shift_scan, KEYEVENTF_KEYUP, None)
        time.sleep(0.04)


def set_clipboard_text(text: str) -> None:
    data = (text + "\0").encode("utf-16-le")
    handle = kernel32.GlobalAlloc(GMEM_MOVEABLE, len(data))
    if not handle:
        raise ctypes.WinError(ctypes.get_last_error())
    locked = kernel32.GlobalLock(handle)
    if not locked:
        raise ctypes.WinError(ctypes.get_last_error())
    ctypes.memmove(locked, data, len(data))
    kernel32.GlobalUnlock(handle)

    if not user32.OpenClipboard(None):
        raise ctypes.WinError(ctypes.get_last_error())
    try:
        user32.EmptyClipboard()
        if not user32.SetClipboardData(CF_UNICODETEXT, handle):
            raise ctypes.WinError(ctypes.get_last_error())
    finally:
        user32.CloseClipboard()


def paste_text(text: str) -> None:
    set_clipboard_text(text)
    press_hotkey(["ctrl", "v"])
    time.sleep(0.5)


def add_months(value: datetime, months: int) -> datetime:
    month_index = value.month - 1 + months
    year = value.year + month_index // 12
    month = month_index % 12 + 1
    first_next_month = datetime(year + (month // 12), (month % 12) + 1, 1)
    last_day = (first_next_month - timedelta(days=1)).day
    return value.replace(year=year, month=month, day=min(value.day, last_day))


def month_datetime(months_offset: int, which: str) -> datetime:
    target = add_months(datetime.now(), months_offset)
    if which == "start":
        target = target.replace(day=1)
    elif which == "mid":
        target = target.replace(day=15)
    elif which == "end":
        target = add_months(target.replace(day=1), 1) - timedelta(days=1)
    else:
        raise ValueError(f"Pilihan tanggal bulan tidak dikenal: {which}")
    return target


def month_date(months_offset: int, which: str, date_format: str) -> str:
    return month_datetime(months_offset, which).strftime(date_format)


def delivery_period_datetime(periods_offset: int, boundary: str) -> datetime:
    today = datetime.now()
    current_half = 0 if today.day <= 15 else 1
    current_index = ((today.year * 12) + (today.month - 1)) * 2 + current_half
    target_index = current_index + periods_offset
    target_month_index, target_half = divmod(target_index, 2)
    year, month_zero_based = divmod(target_month_index, 12)
    month = month_zero_based + 1

    if boundary == "start":
        day = 1 if target_half == 0 else 16
    elif boundary == "end":
        if target_half == 0:
            day = 15
        else:
            first_of_month = datetime(year, month, 1)
            day = (add_months(first_of_month, 1) - timedelta(days=1)).day
    else:
        raise ValueError(f"Boundary periode delivery tidak dikenal: {boundary}")

    return datetime(year, month, day)


def delivery_period_date(periods_offset: int, boundary: str, date_format: str) -> str:
    return delivery_period_datetime(periods_offset, boundary).strftime(date_format)


def type_datetime_picker(dt: datetime) -> None:
    day_str = f"{dt.day:02d}"
    month_str = f"{dt.month:02d}"
    year_str = f"{dt.year:04d}"

    for _ in range(3):
        press_key("left")
        time.sleep(0.04)

    type_text(day_str)
    time.sleep(0.08)

    press_key("right")
    time.sleep(0.08)

    type_text(month_str)
    time.sleep(0.08)

    press_key("right")
    time.sleep(0.08)

    type_text(year_str)
    time.sleep(0.08)


def format_filename(template: str) -> str:
    now = datetime.now()
    return now.strftime(template)


def wait_for_file_stable(path: Path, timeout_seconds: int, stable_seconds: float, since: float = 0.0) -> None:
    deadline = time.time() + timeout_seconds
    last_state: tuple[int, float] | None = None
    stable_since: float | None = None

    while time.time() < deadline:
        if path.exists() and path.is_file():
            stat = path.stat()
            if since and stat.st_mtime < since:
                time.sleep(0.5)
                continue
            state = (stat.st_size, stat.st_mtime)
            if state == last_state:
                if stable_since is None:
                    stable_since = time.time()
                if time.time() - stable_since >= stable_seconds:
                    logging.info("File tersimpan dan stabil: %s", path)
                    return
            else:
                last_state = state
                stable_since = None
        time.sleep(0.5)

    raise TimeoutError(f"File belum selesai tersimpan dalam {timeout_seconds} detik: {path}")



def wait_step(step: dict) -> None:
    seconds = float(step.get("seconds", 1))
    popup_config = step.get("handle_popup")
    until_window = step.get("until_window")
    exit_on_handled = step.get("exit_on_handled", True) if popup_config else False

    # Konfigurasi screen-stable setelah popup ditangani (agar tidak keluar
    # terlalu cepat saat data masih loading)
    post_handle_stable_seconds = float(step.get("post_handle_stable_seconds", 0))
    post_handle_stable_timeout = float(step.get("post_handle_stable_timeout", 0))
    post_handle_stable_threshold = float(step.get("post_handle_stable_threshold", 3.0))
    post_handle_stable_interval = float(step.get("post_handle_stable_interval", 2.0))

    if not popup_config and not until_window:
        time.sleep(seconds)
        return

    ignored_hwnd = user32.GetForegroundWindow()
    deadline = time.time() + seconds
    check_interval = float(popup_config.get("check_interval_seconds", 1)) if popup_config else 0.5

    while time.time() < deadline:
        if until_window:
            found = find_window(until_window)
            if found:
                logging.info("Window target wait step muncul: %s", found[1])
                return

        if popup_config:
            handled = handle_optional_popup(popup_config, ignored_hwnd)
            if handled and exit_on_handled:
                logging.info("Popup berhasil ditangani, keluar dari wait step lebih awal.")
                time.sleep(float(step.get("settle_seconds", 1)))

                # Jika dikonfigurasi, tunggu layar stabil sebelum keluar
                # sehingga data yang masih loading tidak diabaikan
                if post_handle_stable_seconds > 0 and post_handle_stable_timeout > 0:
                    logging.info(
                        "Menunggu layar stabil setelah popup (stable=%.1fs, timeout=%.1fs)...",
                        post_handle_stable_seconds,
                        post_handle_stable_timeout,
                    )
                    stable = wait_for_screen_stable(
                        timeout_seconds=int(post_handle_stable_timeout),
                        stable_seconds=post_handle_stable_seconds,
                        sample_interval=post_handle_stable_interval,
                        diff_threshold=post_handle_stable_threshold,
                        title_contains=step.get("window_title_contains"),
                    )
                    if stable:
                        logging.info("Layar stabil setelah popup, melanjutkan ke step berikutnya.")
                    else:
                        logging.warning(
                            "Layar belum stabil dalam %.1f detik setelah popup, tetap melanjutkan.",
                            post_handle_stable_timeout,
                        )
                return

        time.sleep(check_interval)


def close_excel_processes_and_windows() -> None:
    """Tutup window dan proses Excel/WPS yang terbuka setelah ekspor JO."""
    logging.info("Memastikan aplikasi Excel/WPS ditutup...")
    excel_titles = [
        "JO.xlsx",
        "JO",
        "WPS Office",
        "WPS Spreadsheets",
        "Spreadsheets",
        "Microsoft Excel",
        "Excel",
        "Worksheet in",
        "Worksheet",
        "Book",
    ]

    # 1. Coba tutup window yang terbuka dengan WM_CLOSE dan Alt+F4
    for _ in range(3):
        found = find_window(excel_titles)
        if not found:
            break
        hwnd, title = found
        logging.info("Menutup window Excel/WPS: %s (hwnd=%s)", title, hwnd)
        try:
            focus_window(hwnd)
            time.sleep(0.2)
            user32.PostMessageW(hwnd, WM_CLOSE, 0, 0)
            time.sleep(0.5)
            # Tangani jika ada dialog konfirmasi tutup / simpan
            save_prompt = find_window(["Save", "Simpan", "Confirm", "Prompt", "Microsoft Excel", "WPS"])
            if save_prompt and save_prompt[0] != hwnd:
                prompt_hwnd, prompt_title = save_prompt
                if any(w in prompt_title.lower() for w in ["save", "simpan", "confirm", "prompt"]):
                    logging.info("Menangani dialog konfirmasi tutup Excel: %s", prompt_title)
                    focus_window(prompt_hwnd)
                    press_key("n")  # Don't save
                    time.sleep(0.3)
                    press_key("enter")
            if user32.IsWindowVisible(hwnd):
                press_hotkey(["alt", "f4"])
                time.sleep(0.5)
        except Exception as exc:
            logging.warning("Gagal mengirim sinyal close ke window Excel: %s", exc)

    time.sleep(0.8)

    # 2. Pastikan proses background Excel/WPS di-terminate agar file JO.xlsx tidak terkunci
    excel_procs = ["et.exe", "wps.exe", "excel.exe"]
    for proc_name in excel_procs:
        try:
            res = subprocess.run(
                ["taskkill", "/F", "/IM", proc_name, "/T"],
                capture_output=True,
                text=True,
            )
            if res.returncode == 0:
                logging.info("Proses %s berhasil dibersihkan via taskkill.", proc_name)
        except Exception as exc:
            logging.warning("Gagal taskkill %s: %s", proc_name, exc)


def save_as_file(step: dict) -> Path | None:
    directory = project_path(step.get("directory", ROOT / "downloads"))
    directory.mkdir(parents=True, exist_ok=True)
    filename = format_filename(str(step.get("filename", "jo-export-%Y%m%d-%H%M%S.xlsx")))
    full_path = directory / filename
    fallback_downloads = project_path(step.get("directory", ROOT / "downloads"))
    existing_before = {path.resolve() for path in fallback_downloads.glob("*.xlsx")}

    found = None
    if step.get("force_ctrl_s", False):
        logging.info("Kirim Ctrl+S untuk membuka dialog simpan.")
        press_hotkey(["ctrl", "s"])
        time.sleep(float(step.get("after_ctrl_s_seconds", 1)))
    elif step.get("assume_open", False):
        configured_titles = step.get("title_contains", [])
        if not isinstance(configured_titles, list):
            configured_titles = [configured_titles]
        title_candidates = configured_titles + [
            "Save As", "Simpan Sebagai", "Save", "Simpan",
            "WPS Office", "WPS Spreadsheets", "Spreadsheets", "Microsoft Excel", "Excel", "Book", "JO"
        ]
        title_candidates = list(dict.fromkeys([title for title in title_candidates if title]))

        # Dialog sudah dibuka oleh langkah sebelumnya (misal F12), beri jeda singkat
        time.sleep(0.5)

        # 1. Coba cari title spesifik Save As terlebih dahulu
        found = wait_for_any_window(["Save As", "Simpan Sebagai", "Save", "Simpan"], timeout_seconds=1)

        # 2. Jika tidak ditemukan lewat title spesifik, gunakan foreground window saat ini
        if not found:
            fg_hwnd = user32.GetForegroundWindow()
            if fg_hwnd and user32.IsWindow(fg_hwnd):
                fg_title = window_title(fg_hwnd).strip()
                if not title_matches(fg_title, ["ios-aps", "login", "frmwelcome"]):
                    found = (fg_hwnd, fg_title or "Save As (Foreground)")
                    logging.info("Dialog Save As menggunakan active foreground window: %r (hwnd=%s)", fg_title, fg_hwnd)

        # 3. Jika belum ditemukan, tunggu kandidat window dengan timeout singkat (maks 5 detik)
        if not found:
            timeout_seconds = min(int(step.get("timeout_seconds", 5)), 5)
            found = wait_for_any_window(title_candidates, timeout_seconds)
    else:
        configured_titles = step.get("title_contains", [])
        if not isinstance(configured_titles, list):
            configured_titles = [configured_titles]
        title_contains = configured_titles + ["Save As", "Simpan Sebagai", "Save", "Save File", "Export", "Simpan"]
        title_contains = list(dict.fromkeys([title for title in title_contains if title]))
        timeout_seconds = int(step.get("timeout_seconds", 10))
        found = wait_for_any_window(title_contains, timeout_seconds)
        if not found and step.get("force_ctrl_s_if_not_found", False):
            logging.warning("Popup Save As tidak ditemukan. Kirim Ctrl+S untuk memunculkan dialog simpan.")
            press_hotkey(["ctrl", "s"])
            time.sleep(1)
            found = wait_for_any_window(title_contains, timeout_seconds)
        if not found:
            fg_hwnd = user32.GetForegroundWindow()
            if fg_hwnd and user32.IsWindow(fg_hwnd):
                fg_title = window_title(fg_hwnd).strip()
                if not title_matches(fg_title, ["ios-aps", "login", "frmwelcome"]):
                    found = (fg_hwnd, fg_title or "Save As (Foreground)")

    if not found:
        newest_candidate = None
        newest_mtime = 0.0
        for candidate in fallback_downloads.glob("*.xlsx"):
            try:
                resolved = candidate.resolve()
                stat = candidate.stat()
            except OSError:
                continue
            if resolved in existing_before:
                continue
            if stat.st_mtime > newest_mtime:
                newest_candidate = candidate
                newest_mtime = stat.st_mtime

        if newest_candidate and newest_candidate.exists():
            if full_path.exists():
                try:
                    full_path.unlink()
                except PermissionError:
                    pass
            shutil.copy2(newest_candidate, full_path)
            logging.info("File tersimpan dari fallback download: %s", full_path)
            if step.get("close_after_save", True):
                time.sleep(1.0)
                close_excel_processes_and_windows()
            return full_path

        if step.get("optional", False) or step.get("only_if_window") or step.get("only_if_present", False):
            logging.info("Dialog Save As tidak muncul, melewati penanganan Save As.")
            return None
        raise TimeoutError(f"Dialog Save As tidak muncul dalam {step.get('timeout_seconds', 10)} detik.")

    hwnd, title = found
    logging.info("Popup Save As ditemukan: %s", title)
    focus_window(hwnd)

    # Tangani jika ada dialog error sisa (seperti 'Invalid address')
    invalid_popup = wait_for_any_window(["Invalid address", "Error", "Warning"], 1)
    if invalid_popup:
        logging.info("Menutup popup error Save As sisa: %s", invalid_popup[1])
        press_key("enter")
        time.sleep(0.5)

    # Bersihkan isi field File Name (Ctrl+A lalu Backspace)
    # JANGAN gunakan Alt+N karena di WPS Office tombol 'N' terketik ke field sehingga menjadi 'nE:\...'
    press_hotkey(["ctrl", "a"])
    time.sleep(0.2)
    press_key("backspace")
    time.sleep(0.2)

    save_attempt_at = time.time() - 2.0

    if full_path.exists():
        try:
            full_path.unlink()
        except PermissionError:
            pass

    paste_text(str(full_path.resolve()))
    time.sleep(0.6)
    press_key("enter")
    time.sleep(1.0)

    # 1. Konfirmasi jika window dialog replace file muncul (WPS / Microsoft Excel)
    replace_window_titles = [
        "Duplicate File Already Exists",
        "Duplicate File",
        "Duplicate",
        "Already Exists",
        "Confirm Save As",
        "Confirm",
        "Replace File",
        "Replace",
        "Simpan Sebagai",
        "Konfirmasi",
    ]
    confirm_popup = wait_for_any_window(replace_window_titles, 2)
    if confirm_popup:
        logging.info("Konfirmasi replace file terdeteksi via window: %s (hwnd=%s)", confirm_popup[1], confirm_popup[0])
        focus_window(confirm_popup[0])
        time.sleep(0.3)
        # WPS: default button = [Replace File], hotkey Alt+R
        # Excel: button = [Yes], hotkey Alt+Y
        press_hotkey(["alt", "r"])
        time.sleep(0.2)
        press_hotkey(["alt", "y"])
        time.sleep(0.2)
        press_key("enter")
        time.sleep(0.2)
        press_key("space")
        time.sleep(0.5)

    # 2. Tangani modal WPS internal ("Duplicate File Already Exists") yang berada di dalam Save As window
    # Pada WPS Office modern, dialog duplicate file sering berupa in-window child modal yang menahan Save As.
    for _ in range(4):
        if full_path.exists() and full_path.stat().st_mtime >= save_attempt_at:
            break
        fg_hwnd = user32.GetForegroundWindow()
        if fg_hwnd and user32.IsWindow(fg_hwnd):
            fg_title = window_title(fg_hwnd).strip()
            # Jika window Save As / WPS masih aktif di foreground dan file belum tersimpan
            if fg_hwnd == hwnd or any(t in fg_title.lower() for t in ["save as", "simpan", "wps", "spreadsheets", "excel", "book", "duplicate", "replace", ""]):
                logging.info(
                    "Modal duplicate atau jendela Save As masih aktif (hwnd=%s, title=%r). Mengirim konfirmasi Replace File (Alt+R, Enter, Space)...",
                    fg_hwnd,
                    fg_title,
                )
                focus_window(fg_hwnd)
                time.sleep(0.2)
                press_hotkey(["alt", "r"])
                time.sleep(0.2)
                press_hotkey(["alt", "y"])
                time.sleep(0.2)
                press_key("enter")
                time.sleep(0.2)
                press_key("space")
                time.sleep(0.8)

    try:
        wait_for_file_stable(
            full_path,
            int(step.get("save_timeout_seconds", 180)),
            float(step.get("stable_seconds", 2)),
            since=save_attempt_at,
        )
        logging.info("File tersimpan: %s", full_path)
        if step.get("close_after_save", True):
            time.sleep(1.0)
            close_excel_processes_and_windows()
        return full_path
    except TimeoutError:
        logging.warning("Save dialog tidak menghasilkan file temp. Cek file baru di folder downloads.")

    newest_candidate = None
    newest_mtime = 0.0
    for candidate in fallback_downloads.glob("*.xlsx"):
        try:
            resolved = candidate.resolve()
            stat = candidate.stat()
        except OSError:
            continue
        if resolved in existing_before:
            continue
        if stat.st_mtime > newest_mtime:
            newest_candidate = candidate
            newest_mtime = stat.st_mtime

    if newest_candidate and newest_candidate.exists():
        if full_path.exists():
            try:
                full_path.unlink()
            except PermissionError:
                pass
        shutil.copy2(newest_candidate, full_path)
        logging.info("File tersimpan dari fallback download: %s", full_path)
        if step.get("close_after_save", True):
            time.sleep(1.0)
            close_excel_processes_and_windows()
        return full_path

    if step.get("optional", False) or step.get("only_if_window") or step.get("only_if_present", False):
        logging.info("File Save As belum selesai dan step bersifat opsional, melanjutkan.")
        return None

    raise TimeoutError(f"File hasil export tidak ditemukan di folder downloads: {directory}")


def merge_jo_workbooks(part1_path: Path, part2_path: Path, output_path: Path, archive_parts: bool = True) -> bool:
    """Gabungkan dua file export JO (part1 + part2) menjadi satu file utuh JO.xlsx."""
    import openpyxl

    if not part1_path.exists():
        logging.error("File part 1 tidak ditemukan untuk di-merge: %s", part1_path)
        return False

    if not part2_path.exists():
        logging.warning("File part 2 tidak ditemukan (%s). Menggunakan part 1 sebagai output.", part2_path)
        if output_path.exists():
            try:
                output_path.unlink()
            except Exception:
                pass
        shutil.copy2(part1_path, output_path)
        return True

    logging.info("Memulai penggabungan: %s + %s -> %s", part1_path.name, part2_path.name, output_path.name)
    wb1 = openpyxl.load_workbook(part1_path)
    ws1 = wb1.active

    # 1. Hapus baris 'Total' di akhir ws1 jika ada
    while ws1.max_row > 1:
        first_cell = ws1.cell(row=ws1.max_row, column=1).value
        second_cell = ws1.cell(row=ws1.max_row, column=2).value
        text = str(first_cell or second_cell or "").strip().lower()
        if text == "total":
            logging.info("Menghapus baris Total dari part 1 (row %d)", ws1.max_row)
            ws1.delete_rows(ws1.max_row)
            break
        elif first_cell is None and second_cell is None:
            ws1.delete_rows(ws1.max_row)
        else:
            break

    # 2. Buka ws2 untuk membaca baris data
    wb2 = openpyxl.load_workbook(part2_path, data_only=True)
    ws2 = wb2.active

    # Cari baris header di part 2 (Customer Name / Order No.)
    header_row_idx = None
    for r in range(1, min(15, ws2.max_row + 1)):
        row_vals = [str(ws2.cell(row=r, column=c).value or "").strip().lower() for c in range(1, 10)]
        if any("customer name" in v for v in row_vals) or any("order no" in v for v in row_vals):
            header_row_idx = r
            break

    start_row = (header_row_idx + 1) if header_row_idx else 5
    logging.info("Baris data part 2 dimulai dari baris %d (header baris %s)", start_row, header_row_idx)

    rows_appended = 0
    for r in range(start_row, ws2.max_row + 1):
        first_val = ws2.cell(row=r, column=1).value
        second_val = ws2.cell(row=r, column=2).value
        text = str(first_val or second_val or "").strip().lower()
        if text == "total":
            logging.info("Baris Total part 2 terdeteksi di row %d, pembacaan selesai.", r)
            break

        row_values = [ws2.cell(row=r, column=c).value for c in range(1, ws2.max_column + 1)]
        if any(v is not None and str(v).strip() != "" for v in row_values):
            ws1.append(row_values)
            rows_appended += 1

    logging.info("Berhasil menambahkan %d baris data dari %s.", rows_appended, part2_path.name)

    # 3. Simpan workbook gabungan
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        try:
            output_path.unlink()
        except Exception:
            pass
    wb1.save(output_path)
    wb1.close()
    wb2.close()
    logging.info("File gabungan berhasil disimpan: %s (Total baris: %d)", output_path, ws1.max_row)

    # 4. Arsipkan file part
    if archive_parts:
        archive_dir = output_path.parent.parent / "archive"
        archive_dir.mkdir(parents=True, exist_ok=True)
        now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        for p in [part1_path, part2_path]:
            if p.exists():
                try:
                    shutil.move(str(p), str(archive_dir / f"{p.stem}_{now_str}{p.suffix}"))
                    logging.info("File part %s diarsipkan ke %s", p.name, archive_dir.name)
                except Exception as exc:
                    logging.warning("Gagal mengarsipkan %s: %s", p.name, exc)
                    try:
                        p.unlink()
                    except Exception:
                        pass

    return True


def merge_excel_parts_step(step: dict) -> None:
    directory = project_path(step.get("directory", ROOT / "downloads"))
    part1_name = str(step.get("part1", "JO_part1.xlsx"))
    part2_name = str(step.get("part2", "JO_part2.xlsx"))
    target_name = str(step.get("target", "JO.xlsx"))

    part1 = directory / part1_name
    part2 = directory / part2_name
    target = directory / target_name

    merge_jo_workbooks(part1, part2, target, archive_parts=bool(step.get("archive_parts", True)))


def close_aps_windows(config: dict) -> None:
    """Close APS, Excel, and any export dialog after a failed GUI step."""
    try:
        close_excel_processes_and_windows()
    except Exception as exc:
        logging.warning("Gagal menutup Excel saat close_aps_windows: %s", exc)

    title_options = config.get("window_title_contains", ["IOS-APS", "Login"])
    if not isinstance(title_options, list):
        title_options = [title_options]

    for _ in range(5):
        # 1. Dahulukan konfirmasi dialog Query / Are you sure jika muncul
        query = wait_for_any_window(["Query", "Are you sure", "Exit"], 1)
        if query:
            hwnd, title = query
            logging.info("Menutup dialog konfirmasi keluar APS: %s", title)
            focus_window(hwnd)
            press_key("enter")
            time.sleep(1)
            continue

        found = wait_for_any_window(title_options + ["Save As", "Save", "Export"], 1)
        if not found:
            return

        hwnd, title = found
        logging.warning("Menutup window APS: %s", title)
        focus_window(hwnd)
        press_hotkey(["alt", "f4"])
        time.sleep(0.8)
        query = wait_for_any_window(["Query", "Are you sure", "Exit"], 2)
        if query:
            focus_window(query[0])
            press_key("enter")
        time.sleep(1)


def screenshot(name: str) -> Path | None:
    if ImageGrab is None:
        logging.warning("Pillow ImageGrab tidak tersedia, screenshot dilewati.")
        return None
    screenshot_dir = ROOT / "screenshots"
    screenshot_dir.mkdir(exist_ok=True)
    path = screenshot_dir / f"{datetime.now():%Y%m%d-%H%M%S}-{name}.png"
    try:
        ImageGrab.grab().save(path)
        logging.info("Screenshot: %s", path)
    except OSError as error:
        logging.warning("Screenshot dilewati: %s", error)
        return None
    return path


def window_bbox(hwnd: int) -> tuple[int, int, int, int] | None:
    if not hwnd:
        return None

    rect = RECT()
    if not user32.GetWindowRect(hwnd, ctypes.byref(rect)):
        return None

    if rect.right <= rect.left or rect.bottom <= rect.top:
        return None

    return rect.left, rect.top, rect.right, rect.bottom


def wait_for_screen_stable(timeout_seconds: int, stable_seconds: float = 8.0, sample_interval: float = 2.0, diff_threshold: float = 3.0, title_contains=None) -> bool:
    if ImageGrab is None or ImageChops is None or ImageStat is None:
        logging.warning("Pillow tidak tersedia, fallback ke wait biasa.")
        time.sleep(timeout_seconds)
        return False

    target_hwnd = None
    if title_contains:
        try:
            target_hwnd, title = wait_for_window(title_contains, min(10, timeout_seconds))
            logging.info("Stabilisasi dipantau pada window: %s", title)
        except TimeoutError:
            logging.warning("Window target stabilisasi tidak ditemukan, fallback ke foreground window.")

    deadline = time.time() + timeout_seconds
    stable_since = None
    previous = None

    while time.time() < deadline:
        bbox = window_bbox(target_hwnd) if target_hwnd else window_bbox(user32.GetForegroundWindow())
        if not bbox:
            time.sleep(sample_interval)
            continue

        try:
            current = ImageGrab.grab(bbox=bbox).convert("L")
        except OSError as error:
            logging.warning("Gagal ambil screenshot untuk cek stabilitas: %s", error)
            time.sleep(sample_interval)
            continue

        if previous is not None:
            diff = ImageChops.difference(previous, current)
            stat = ImageStat.Stat(diff)
            mean_diff = stat.mean[0] if stat.mean else 0.0
            if mean_diff <= diff_threshold:
                if stable_since is None:
                    stable_since = time.time()
                if time.time() - stable_since >= stable_seconds:
                    logging.info("Layar APS sudah stabil.")
                    return True
            else:
                stable_since = None
        previous = current
        time.sleep(sample_interval)

    logging.warning("Layar APS belum stabil dalam %s detik.", timeout_seconds)
    return False


def project_path(value) -> Path:
    path = Path(str(value))
    return path if path.is_absolute() else ROOT / path


def archive_stale_download(download_path: Path) -> bool:
    if not download_path.is_file():
        return False

    file_date = datetime.fromtimestamp(download_path.stat().st_mtime).date()
    today = datetime.now().date()
    if file_date >= today:
        return False

    for attempt in range(1, 6):
        try:
            archive_dir = ARCHIVE_DIR / file_date.strftime("%Y-%m")
            archive_dir.mkdir(parents=True, exist_ok=True)
            archive_path = archive_dir / f"{file_date:%Y-%m-%d}_{download_path.name}"
            if archive_path.exists():
                archive_path.unlink()
            shutil.move(str(download_path), str(archive_path))
            logging.info("File diarsipkan: %s", archive_path)
            return True
        except PermissionError:
            if attempt == 5:
                logging.warning("Gagal memindah file lama ke archive karena masih dipakai: %s", download_path)
                return False
            time.sleep(2)

    return False


def archive_download_folder(source_dir: Path, patterns: list[str]) -> list[Path]:
    if not source_dir.exists():
        logging.warning("Folder download browser tidak ditemukan: %s", source_dir)
        return []

    archived: list[Path] = []
    for pattern in patterns:
        for path in source_dir.glob(pattern):
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".xlsx", ".xls", ".html"}:
                continue
            if archive_stale_download(path):
                archived.append(path)
    return archived


def wait_for_downloads(source_dir: Path, patterns: list[str], since: float, timeout_seconds: int) -> list[Path]:
    if not source_dir.exists():
        logging.warning("Folder download browser tidak ditemukan: %s", source_dir)
        return []

    deadline = time.time() + timeout_seconds
    temporary_suffixes = (".crdownload", ".download", ".part", ".tmp")
    while time.time() < deadline:
        matches: list[Path] = []
        temporary_files = []
        for pattern in patterns:
            for source in source_dir.glob(pattern):
                if source.is_file() and source.stat().st_mtime >= since:
                    matches.append(source)
        for source in source_dir.iterdir():
            if source.is_file() and source.suffix.lower() in temporary_suffixes and source.stat().st_mtime >= since:
                temporary_files.append(source)
        if matches and not temporary_files:
            return matches
        time.sleep(1)

    logging.warning("Tidak ada file download baru yang selesai dalam %s detik.", timeout_seconds)
    return []


def run_step(step: dict) -> None:
    action = step.get("action")
    log_step = dict(step)
    if "password" in str(log_step).lower() or log_step.get("env"):
        log_step = {**log_step, "text": "***"}
    logging.info("Step: %s", log_step)
    if action == "wait":
        wait_step(step)
    elif action == "wait_for_screen_stable":
        wait_for_screen_stable(
            int(step.get("timeout_seconds", 120)),
            float(step.get("stable_seconds", 8)),
            float(step.get("sample_interval", 2)),
            float(step.get("diff_threshold", 3)),
            step.get("window_title_contains"),
        )
    elif action == "key":
        for _ in range(int(step.get("count", 1))):
            press_key(step["key"])
            if step.get("interval"):
                time.sleep(float(step.get("interval")))
        if step.get("after_seconds"):
            time.sleep(float(step.get("after_seconds")))
    elif action == "hotkey":
        press_hotkey(step["keys"])
        if step.get("after_seconds"):
            time.sleep(float(step.get("after_seconds")))
        if "excel" in step.get("name", "").lower():
            time.sleep(0.5)
            close_excel_processes_and_windows()
    elif action == "close_excel":
        close_excel_processes_and_windows()
    elif action == "click":
        for _ in range(int(step.get("count", 1))):
            fg_hwnd = user32.GetForegroundWindow()
            active_title = window_title(fg_hwnd).strip()
            x = int(step["x"])
            y = int(step["y"])
            if step.get("window_relative", False):
                bbox = window_bbox(fg_hwnd)
                if bbox:
                    x = bbox[0] + x
                    y = bbox[1] + y
                elif step.get("fallback_absolute_x") and step.get("fallback_absolute_y"):
                    x = int(step["fallback_absolute_x"])
                    y = int(step["fallback_absolute_y"])
            logging.info("Klik koordinat x=%s y=%s pada window aktif: %s", x, y, active_title)
            click(x, y)
            time.sleep(float(step.get("interval", 0.15)))
    elif action == "type":
        type_text(str(step.get("text", "")))
    elif action == "paste":
        paste_text(str(step.get("text", "")))
    elif action == "type_month_date":
        dt = month_datetime(
            int(step.get("months_offset", 0)),
            str(step.get("which", "start")),
        )
        logging.info("Ketik tanggal bulan (DateTimePicker): %s", dt.strftime("%d/%m/%Y"))
        type_datetime_picker(dt)
    elif action == "paste_month_date":
        text = month_date(
            int(step.get("months_offset", 0)),
            str(step.get("which", "start")),
            str(step.get("format", "%d/%m/%Y")),
        )
        logging.info("Paste tanggal bulan: %s", text)
        paste_text(text)
    elif action == "type_delivery_date":
        dt = delivery_period_datetime(
            int(step.get("periods_offset", 0)),
            str(step.get("boundary", "start")),
        )
        logging.info("Ketik tanggal delivery (DateTimePicker): %s", dt.strftime("%d/%m/%Y"))
        type_datetime_picker(dt)
    elif action == "paste_delivery_date":
        text = delivery_period_date(
            int(step.get("periods_offset", 0)),
            str(step.get("boundary", "start")),
            str(step.get("format", "%d/%m/%Y")),
        )
        logging.info("Paste tanggal delivery: %s", text)
        paste_text(text)
    elif action == "save_as":
        save_as_file(step)
    elif action == "merge_excel_parts":
        merge_excel_parts_step(step)
    elif action == "type_env":
        name = str(step["env"])
        value = os.environ.get(name)
        if value is None:
            env_hint = f"Buat file {DEFAULT_ENV} atau set di PowerShell sebelum menjalankan python main.py."
            if (ROOT / ".env.example").exists() and not DEFAULT_ENV.exists():
                env_hint = f"File .env belum ada. Salin .env.example menjadi {DEFAULT_ENV}, lalu isi credential."
            raise ValueError(f"Environment variable belum diset: {name}. {env_hint}")
        type_text(value)
    elif action == "screenshot":
        screenshot(str(step.get("name", "step")))
    else:
        raise ValueError(f"Action tidak dikenal: {action}")


def focus_step_window(step: dict, default_hwnd: int) -> int:
    if step.get("action") in ("close_excel", "merge_excel_parts"):
        return default_hwnd
    title_contains = step.get("window_title_contains")
    if title_contains:
        options = title_contains if isinstance(title_contains, list) else [title_contains]
        if any("login" in str(t).lower() for t in options):
            logging.info("Menunggu splash screen FrmWelcomeLoading selesai sebelum fokus ke Login...")
            wait_window_disappear(["FrmWelcomeLoading"], 30)
            time.sleep(1)

        popup_config = step.get("handle_popup")
        hwnd, title = wait_for_window(title_contains, int(step.get("window_timeout_seconds", 60)), popup_config)
        logging.info("Window target step ditemukan: %s", title)
        should_maximize = step.get("maximize", not title_matches(title, ["login", "error", "warning", "frmwelcomeloading"]))
        focus_window(hwnd, should_maximize)
        return hwnd
    if step.get("keep_current_window", False) or step.get("action") == "wait":
        return default_hwnd
    if user32.GetForegroundWindow() != default_hwnd:
        title = window_title(default_hwnd)
        should_maximize = not title_matches(title, ["login", "error", "warning", "frmwelcomeloading"])
        focus_window(default_hwnd, should_maximize)
    return default_hwnd


def handle_security_warning(config: dict) -> None:
    warning = config.get("security_warning", {})
    if not warning.get("enabled", True):
        return

    title_options = warning.get("title_contains", ["Open File", "Security Warning"])
    timeout = int(warning.get("timeout_seconds", 8))
    found = wait_for_any_window(title_options, timeout)
    if not found:
        logging.info("Popup security warning tidak muncul.")
        return

    hwnd, title = found
    logging.info("Popup security warning ditemukan: %s", title)
    focus_window(hwnd)
    screenshot("security-warning")

    for step in warning.get("steps", [{"action": "hotkey", "keys": ["alt", "r"]}]):
        run_step(step)
        time.sleep(0.5)


def handle_already_running_warning(config: dict) -> None:
    warning = config.get("already_running_warning", {})
    if not warning.get("enabled", True):
        return

    found = wait_for_any_window(warning.get("title_contains", ["Error"]), int(warning.get("timeout_seconds", 2)))
    if not found:
        return

    hwnd, title = found
    logging.info("Popup aplikasi sudah berjalan ditemukan: %s", title)
    focus_window(hwnd)
    screenshot("already-running-warning")
    for step in warning.get("steps", [{"action": "key", "key": "enter"}]):
        run_step(step)
        time.sleep(0.5)


def start_shortcut(shortcut: Path, config: dict) -> None:
    launch_cfg = config.get("launch", {})
    bypass_zone_check = launch_cfg.get("bypass_zone_check", True)

    old_zone_check = os.environ.get("SEE_MASK_NOZONECHECKS")
    if bypass_zone_check:
        logging.info("Bypass Open File security warning aktif untuk proses launch ini.")
        os.environ["SEE_MASK_NOZONECHECKS"] = "1"

    try:
        if shortcut.suffix.lower() == ".exe":
            subprocess.Popen(
                [str(shortcut)],
                cwd=str(shortcut.parent) if shortcut.parent.exists() else None,
                creationflags=getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0),
            )
        else:
            os.startfile(shortcut)
    finally:
        if bypass_zone_check:
            if old_zone_check is None:
                os.environ.pop("SEE_MASK_NOZONECHECKS", None)
            else:
                os.environ["SEE_MASK_NOZONECHECKS"] = old_zone_check


def ensure_excel_file_associations() -> None:
    """Pastikan ekstensi spreadsheet (.xls, .xlsx) terdaftar ke handler spreadsheet yang valid."""
    try:
        import winreg
        for ext, progid in [(".xls", "ET.Xls.6"), (".xlsx", "ET.Xlsx.6")]:
            with winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf"Software\Classes\{ext}") as k:
                winreg.SetValue(k, "", winreg.REG_SZ, progid)
        logging.info("Asosiasi ekstensi spreadsheet (.xls, .xlsx) dipastikan ke WPS (ET.Xls.6, ET.Xlsx.6).")
    except Exception as exc:
        logging.warning("Gagal memastikan asosiasi file spreadsheet di registry HKCU: %s", exc)


def cleanup_stale_processes(config: dict) -> None:
    """Tutup sisa window/proses IOS-APS, Excel, dan file temporary lock sebelum memulai proses baru."""
    logging.info("Memulai pre-run cleanup sisa proses...")
    ensure_excel_file_associations()
    for proc_name in ["IOS-APS.exe", "excel.exe", "wps.exe", "et.exe"]:
        try:
            subprocess.run(["taskkill", "/F", "/IM", proc_name, "/T"], capture_output=True)
        except Exception:
            pass
    time.sleep(1)

    # 1. Bersihkan popup sisa jika ada sebelum mulai baru
    for _ in range(3):
        found = wait_for_any_window(["Query", "Are you sure", "Prompt", "FrmWelcomeLoading", "Error"], 1)
        if found:
            hwnd, title = found
            logging.info("Menutup popup sisa: %s (hwnd=%s)", title, hwnd)
            focus_window(hwnd)
            press_key("enter")
            time.sleep(0.5)
            user32.PostMessageW(hwnd, WM_CLOSE, 0, 0)
            time.sleep(0.5)

    # 2. Tutup window IOS-APS sisa jika ada sebelum mulai baru
    found_aps = find_window(["IOS-APS"])
    if found_aps:
        hwnd, title = found_aps
        logging.warning("Ditemukan window IOS-APS lama: %s. Menutup untuk fresh start...", title)
        focus_window(hwnd)
        press_hotkey(["alt", "f4"])
        time.sleep(1)
        query = wait_for_any_window(["Query", "Are you sure"], 3)
        if query:
            focus_window(query[0])
            press_key("enter")
            time.sleep(1)

    # 3. Tutup window Login sisa jika ada dari run gagal sebelumnya
    found_login = find_window(["Login"])
    if found_login:
        hwnd, title = found_login
        logging.warning("Ditemukan window Login sisa: %s. Menutup untuk fresh start...", title)
        user32.PostMessageW(hwnd, WM_CLOSE, 0, 0)
        time.sleep(0.5)

    # 4. Bersihkan temporary lock file Excel dan backup JO.xlsx lama di folder downloads
    downloads_dir = ROOT / "downloads"
    if downloads_dir.exists():
        for lock_file in downloads_dir.glob("~$*"):
            try:
                lock_file.unlink()
                logging.info("File temporary lock dibersihkan: %s", lock_file.name)
            except Exception:
                pass

        # Bersihkan file part lama jika ada sisa run sebelumnya
        for part_file in downloads_dir.glob("JO_part*.xlsx"):
            try:
                part_file.unlink()
                logging.info("File part lama dibersihkan: %s", part_file.name)
            except Exception:
                pass

        # Backup file JO.xlsx lama ke archive agar saat Save As di WPS tidak terjadi konflik "Duplicate File Already Exists"
        old_jo = downloads_dir / "JO.xlsx"
        if old_jo.exists():
            try:
                archive_dir = ROOT / "archive"
                archive_dir.mkdir(parents=True, exist_ok=True)
                mtime = datetime.fromtimestamp(old_jo.stat().st_mtime)
                backup_name = f"JO_backup_{mtime.strftime('%Y%m%d_%H%M%S')}.xlsx"
                shutil.copy2(old_jo, archive_dir / backup_name)
                old_jo.unlink()
                logging.info("File JO.xlsx lama dibackup ke archive (%s) dan dibersihkan dari downloads.", backup_name)
            except Exception as exc:
                logging.warning("Gagal membersihkan JO.xlsx lama saat cleanup: %s", exc)


def main() -> int:
    parser = argparse.ArgumentParser(description="RPA download IOS-APS")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="Path config JSON")
    parser.add_argument("--check-only", action="store_true", help="Validasi config tanpa membuka aplikasi")
    parser.add_argument("--list-windows", action="store_true", help="Tampilkan daftar window aktif lalu keluar")
    args = parser.parse_args()

    log_path = setup_logging()
    load_env_file(DEFAULT_ENV)
    config = load_config(Path(args.config))
    shortcut = Path(config["shortcut_path"])
    if not shortcut.exists():
        raise FileNotFoundError(f"Shortcut tidak ditemukan: {shortcut}")

    logging.info("Log: %s", log_path)
    if args.list_windows:
        for hwnd, title in visible_windows():
            logging.info("Window aktif: %s | %s", hwnd, title)
        return 0
    if args.check_only:
        logging.info("Config valid. Shortcut ditemukan: %s", shortcut)
        logging.info("Mode check-only aktif, aplikasi tidak dibuka.")
        return 0

    download_started_at = time.time()

    cleanup_stale_processes(config)

    logging.info("Membuka shortcut: %s", shortcut)
    start_shortcut(shortcut, config)
    handle_security_warning(config)
    handle_already_running_warning(config)

    hwnd, title = wait_for_window(config.get("window_title_contains"), int(config.get("window_timeout_seconds", 60)))
    logging.info("Window ditemukan: %s", title)
    if "frmwelcomeloading" in title.lower():
        logging.info("Splash screen FrmWelcomeLoading terdeteksi, menunggu selesai...")
        wait_window_disappear(["FrmWelcomeLoading"], 30)
        time.sleep(1)
    should_maximize = not title_matches(title, ["login", "error", "warning", "frmwelcomeloading"])
    focus_window(hwnd, should_maximize)
    screenshot("opened")

    try:
        current_hwnd = hwnd
        for step in config.get("steps", []):
            try:
                only_if_win = step.get("only_if_window")
                if only_if_win:
                    win_targets = only_if_win if isinstance(only_if_win, list) else [only_if_win]
                    timeout = int(step.get("window_timeout_seconds", 5))
                    found_win = wait_for_any_window(win_targets, timeout)
                    if not found_win:
                        logging.info("Window %s tidak muncul, step '%s' dilewati.", win_targets, step.get("name", "unnamed"))
                        continue
                    current_hwnd = found_win[0]
                    should_maximize = step.get("maximize", False)
                    focus_window(current_hwnd, should_maximize)

                current_hwnd = focus_step_window(step, current_hwnd)
                run_step(step)
            except Exception as step_exc:
                if step.get("optional", False):
                    logging.warning("Step opsional dilewati karena gagal: %s (Detail: %s)", step.get("name", "unnamed"), step_exc)
                    continue
                raise
    except Exception:
        logging.exception("RPA gagal pada langkah GUI APS.")
        close_aps_windows(config)
        raise

    # Pastikan Excel/WPS ditutup dan file temporary lock Excel dibersihkan jika masih tersisa
    close_excel_processes_and_windows()
    downloads_dir = ROOT / "downloads"
    if downloads_dir.exists():
        for lock_file in downloads_dir.glob("~$*"):
            try:
                lock_file.unlink()
            except Exception:
                pass

        # Safety fallback: jika JO_part1 ada dan JO.xlsx belum terbuat (misal step merge terlewat), jalankan merge otomatis
        part1_fallback = downloads_dir / "JO_part1.xlsx"
        part2_fallback = downloads_dir / "JO_part2.xlsx"
        final_jo = downloads_dir / "JO.xlsx"
        if part1_fallback.exists() and not final_jo.exists():
            logging.info("Safety fallback: Menjalankan auto-merge untuk JO_part...")
            merge_jo_workbooks(part1_fallback, part2_fallback, final_jo)

    download_cfg = config.get("copy_downloads")
    if download_cfg and download_cfg.get("enabled"):
        source_dir = project_path(download_cfg.get("source_dir", str(Path.home() / "Downloads")))
        archive_dir = project_path(download_cfg.get("archive_dir", str(ROOT / "archive")))
        patterns = download_cfg.get("patterns", ["*"])
        timeout_seconds = int(download_cfg.get("timeout_seconds", 120))
        wait_for_downloads(source_dir, patterns, download_started_at, timeout_seconds)
        archived = archive_download_folder(source_dir, patterns)
        if not archived:
            logging.warning("Tidak ada file lama yang diarsipkan.")

    screenshot("finished")
    logging.info("Selesai.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        logging.exception("RPA gagal.")
        raise
