import glob
import openpyxl

files = glob.glob('rpa/aps-rpa/downloads/*.xlsx') + glob.glob('rpa/aps-rpa/archive/*.xlsx')
if files:
    fpath = files[0]
    print("Reading", fpath)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        print("Sheet:", sheetname, "rows:", sheet.max_row, "cols:", sheet.max_column)
        rows = list(sheet.iter_rows(values_only=True, max_row=5))
        for r in rows:
            if any(r):
                print("  ", [str(c)[:30] if c is not None else '' for c in r[:25]])
else:
    print("No APS xlsx found")

