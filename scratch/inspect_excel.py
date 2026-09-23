import openpyxl
from pathlib import Path

files = [
    Path("rpa/engage-rpa/downloads/32a_engage.xlsx"),
    Path("rpa/engage-rpa/archive/2026-09/2026-09-01_32a_inflow.xlsx"),
    Path("rpa/engage-rpa/archive/2026-09/2026-09-01_32a_outflow.xlsx"),
    Path("rpa/engage-rpa/archive/2026-08/2026-08-28_32_engage.xlsx"),
]

for f in files:
    if not f.is_file():
        print(f"{f} not found")
        continue
    wb = openpyxl.load_workbook(f, read_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True, max_row=10))
    print(f"\n=== {f} ===")
    print(f"Sheet name: {sheet.title}, total rows sample: {len(rows)}")
    header = None
    for r in rows:
        r_str = [str(c).strip() for c in r if c is not None]
        if "Date" in r_str or "We_datum" in r_str:
            header = [str(c).strip() if c is not None else "" for c in r]
            break
    if header:
        print(f"Header: {header[:8]}")
        # find Storage Nr column
        storage_col = None
        for i, h in enumerate(header):
            if h in ("Storage Nr", "We_lagnr"):
                storage_col = i
                break
        date_col = None
        for i, h in enumerate(header):
            if h in ("Date", "We_datum"):
                date_col = i
                break
        print(f"date_col: {date_col}, storage_col: {storage_col}")
        
        # sample next rows
        sample_rows = list(sheet.iter_rows(values_only=True, min_row=3, max_row=15))
        for sr in sample_rows[:5]:
            d = sr[date_col] if date_col is not None and len(sr) > date_col else None
            st = sr[storage_col] if storage_col is not None and len(sr) > storage_col else None
            print(f"  sample row -> date: {d}, storage: {st}")
    else:
        print("No header found")
