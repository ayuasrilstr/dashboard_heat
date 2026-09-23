import openpyxl

def inspect_excel(path):
    print("=== File:", path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        for name in wb.sheetnames:
            sheet = wb[name]
            rows = list(sheet.iter_rows(values_only=True, max_row=5))
            print(f"Sheet: {name}, rows: {sheet.max_row}, cols: {sheet.max_column}")
            for r in rows:
                if any(r):
                    print("  ", [str(c)[:25] if c is not None else '' for c in r[:15]])
    except Exception as e:
        print("Error:", e)

inspect_excel('rpa/accessories-rpa/downloads/CONTROLIST.xlsx')

